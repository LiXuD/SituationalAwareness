#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
I-06 开发期脚本（仅本地运行，用托管 venv 的 openpyxl 执行，
容器运行时【不】依赖此脚本，也不依赖 openpyxl）：
  - 生成 Excel 导入模板  services/asset/templates/asset-import-template.xlsx
  - 生成自测样本          samples/assets/valid-assets.xlsx / invalid-assets.xlsx

表头顺序必须与 services/asset/asset_server.py 中的常量保持一致：
  资产编号(可选), 资产名称, IP地址, 资产类型, 重要度, 风险评分,
  责任人, 所属部门, 位置/机房, 操作系统, 标签, 备注

用法（在 WorkBuddy 沙箱内）：
  export PATH=/usr/local/bin:$PATH
  /Users/lixd/.workbuddy/binaries/python/envs/default/bin/python scripts/gen-asset-template.py
"""
import argparse
import os
import sys

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    sys.stderr.write("需要 openpyxl（请用托管 venv 执行：~/.workbuddy/binaries/python/envs/default/bin/python）\n")
    sys.exit(2)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE_OUT = os.path.join(ROOT, "services", "asset", "templates", "asset-import-template.xlsx")
SAMPLES_DIR = os.path.join(ROOT, "samples", "assets")

HEADERS = ["资产编号(可选)", "资产名称", "IP地址", "资产类型", "重要度", "风险评分",
           "责任人", "所属部门", "位置/机房", "操作系统", "标签", "备注"]

FIELD_DOCS = [
    ("资产编号(可选)", "资产唯一 ID；留空则由系统自动生成。若填写且已存在则按该 ID 更新（upsert）。"),
    ("资产名称", "必填。资产名称/用途简述。"),
    ("IP地址", "必填。IPv4 或 IPv6 地址，格式校验。"),
    ("资产类型", "枚举：服务器/网络设备/安全设备/终端设备/应用系统/数据库/其他；留空默认其他。"),
    ("重要度", "必填。枚举：核心 / 重要 / 一般（对应权重 3/2/1，供大屏热力着色）。"),
    ("风险评分", "必填。整数 0-100。"),
    ("责任人", "选填。资产负责人。"),
    ("所属部门", "选填。"),
    ("位置/机房", "选填。物理位置。"),
    ("操作系统", "选填。"),
    ("标签", "选填。逗号分隔。"),
    ("备注", "选填。"),
]

EXAMPLE_ROW = ["", "核心Web服务器", "10.20.30.40", "服务器", "核心", "90",
               "张伟", "信息技术部", "机房A-01", "Linux-CentOS7", "web,生产", "生产核心资产"]

VALID_ROWS = [
    ["AST-001", "核心Web服务器", "10.20.30.40", "服务器", "核心", "90", "张伟", "信息技术部", "机房A-01", "Linux-CentOS7", "web,生产", "生产核心资产"],
    ["AST-002", "边界防火墙", "203.0.113.1", "安全设备", "核心", "85", "李娜", "网络组", "机房A-02", "专用OS", "fw,边界", "互联网出口防火墙"],
    ["AST-003", "办公终端-笔记本", "192.168.10.55", "终端设备", "一般", "20", "王芳", "综合部", "办公区", "Windows11", "pc", "员工笔记本"],
    ["AST-004", "主数据库", "10.20.30.100", "数据库", "重要", "70", "赵强", "信息技术部", "机房B-01", "Linux-Rocky", "db,核心", "业务主库"],
    ["AST-005", "对外API网关", "10.20.30.200", "应用系统", "重要", "60", "孙磊", "研发组", "机房B-02", "Linux-Ubuntu", "api", "API 网关"],
]

# 含非法行：第2行 IP 非法；第3行 重要度枚举非法；第4行 风险评分 >100；第5行 缺资产名称
INVALID_ROWS = [
    ["BAD-001", "正常资产A", "10.0.0.1", "服务器", "核心", "50", "甲", "部门", "机房", "Linux", "", "合法"],
    ["BAD-002", "IP非法", "999.1.1.1", "服务器", "核心", "50", "乙", "部门", "机房", "Linux", "", "IP 格式错误"],
    ["BAD-003", "重要度非法", "10.0.0.3", "服务器", "极高", "50", "丙", "部门", "机房", "Linux", "", "重要度应为核心/重要/一般"],
    ["BAD-004", "风险超界", "10.0.0.4", "服务器", "一般", "150", "丁", "部门", "机房", "Linux", "", "风险评分应 0-100"],
    ["", "", "10.0.0.5", "服务器", "重要", "40", "戊", "部门", "机房", "Linux", "", "资产名称为空（必填缺失）"],
]


def style_header(ws):
    fill = PatternFill("solid", fgColor="1e2536")
    font = Font(bold=True, color="FFFFFF")
    for c in ws[1]:
        c.fill = fill
        c.font = font
        c.alignment = Alignment(horizontal="center", vertical="center")


def write_sheet(ws, rows):
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    style_header(ws)
    widths = [16, 18, 16, 12, 10, 10, 10, 12, 12, 14, 14, 22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def build_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "资产导入模板"
    write_sheet(ws, [EXAMPLE_ROW])
    doc_ws = wb.create_sheet("字段说明")
    doc_ws.append(["字段", "说明"])
    style_header(doc_ws)
    for name, desc in FIELD_DOCS:
        doc_ws.append([name, desc])
    doc_ws.column_dimensions["A"].width = 18
    doc_ws.column_dimensions["B"].width = 70
    os.makedirs(os.path.dirname(TEMPLATE_OUT), exist_ok=True)
    wb.save(TEMPLATE_OUT)
    print("[*] 模板已生成:", TEMPLATE_OUT)


def build_samples():
    os.makedirs(SAMPLES_DIR, exist_ok=True)
    for kind, rows in (("valid", VALID_ROWS), ("invalid", INVALID_ROWS)):
        wb = Workbook()
        ws = wb.active
        ws.title = "资产导入"
        write_sheet(ws, rows)
        out = os.path.join(SAMPLES_DIR, "%s-assets.xlsx" % kind)
        wb.save(out)
        print("[*] 样本已生成:", out, "(%d 数据行)" % len(rows))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--kind", choices=["template", "samples", "all"], default="all")
    args = ap.parse_args()
    if args.kind in ("template", "all"):
        build_template()
    if args.kind in ("samples", "all"):
        build_samples()


if __name__ == "__main__":
    main()
